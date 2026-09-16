import unittest
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from flat_file_renderers.xlsx import XlsxRenderer


class XlsxRendererTestCase(unittest.TestCase):
    def _read_single_sheet(self, zip_path):
        with ZipFile(zip_path) as zf:
            names = zf.namelist()
            self.assertEqual(len(names), 1)
            wb = load_workbook(BytesIO(zf.read(names[0])))
            ws = wb.active
            return [[cell.value for cell in row] for row in ws.iter_rows()]

    def test_renders_simple_dataset_to_single_sheet(self):
        renderer = XlsxRenderer({})
        context = {"dataset": [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]}
        zip_path = renderer.render(context)
        try:
            rows = self._read_single_sheet(zip_path)
            self.assertEqual(rows[0], ["age", "name"])
            # Values are always written as text (str()), by design - even numeric columns
            # become strings in the cell.
            self.assertEqual(sorted(rows[1:]), [["25", "Bob"], ["30", "Alice"]])
        finally:
            zip_path.unlink(missing_ok=True)

    def test_renders_empty_dataset_with_no_rows(self):
        renderer = XlsxRenderer({})
        zip_path = renderer.render({"dataset": []})
        try:
            rows = self._read_single_sheet(zip_path)
            self.assertEqual(rows, [])
        finally:
            zip_path.unlink(missing_ok=True)

    def test_multi_sheet_dataset_creates_one_sheet_per_key(self):
        renderer = XlsxRenderer({})
        context = {
            "dataset": {
                "Students": [{"name": "Alice"}],
                "Courses": [{"title": "Python"}],
            }
        }
        zip_path = renderer.render(context)
        try:
            with ZipFile(zip_path) as zf:
                names = zf.namelist()
                self.assertEqual(len(names), 1)
                wb = load_workbook(BytesIO(zf.read(names[0])))
                self.assertEqual(wb.sheetnames, ["Students", "Courses"])
                self.assertEqual([cell.value for cell in next(wb["Students"].iter_rows())], ["name"])
        finally:
            zip_path.unlink(missing_ok=True)
